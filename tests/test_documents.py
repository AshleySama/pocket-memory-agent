from __future__ import annotations

import base64
import tempfile
import unittest
from pathlib import Path

import pymupdf
import numpy as np
from docx import Document as WordDocument
from openpyxl import Workbook

from pocket_memory.documents import DocumentExtractor
from pocket_memory.embedding import EmbeddingService
from pocket_memory.rag.pipeline import RagPipeline
from pocket_memory.storage import NoteStore


class DocumentEmbeddingEngine:
    available = True
    unavailable_reason = ""

    def embed(self, text: str):
        if "深层结论" in text or "深层问题" in text:
            return np.asarray([1.0, 0.0], dtype=np.float32)
        return np.asarray([0.0, 1.0], dtype=np.float32)


class DocumentImportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.store = NoteStore(self.root / "data")

    def tearDown(self) -> None:
        self.store.close()
        self.temp.cleanup()

    def create_document(self, name: str, data: bytes):
        data_url = "data:application/octet-stream;base64," + base64.b64encode(data).decode("ascii")
        return self.store.create_document_from_data_url(name, data_url)

    def test_extractors_keep_locations_for_txt_docx_xlsx_and_pdf(self) -> None:
        text_document = self.create_document("说明.txt", "第一段\n\n采购阈值为 50 万元。".encode("utf-8"))
        text_chunks = DocumentExtractor.extract(text_document, self.store.data_dir / text_document.stored_path)
        self.assertEqual(text_chunks[0]["location"], "正文")
        self.assertTrue(any("采购阈值" in chunk["text"] for chunk in text_chunks))

        word_path = self.root / "流程.docx"
        word = WordDocument()
        word.add_heading("审批流程", level=1)
        word.add_paragraph("红色预警需要总监审批。")
        word.save(word_path)
        word_document = self.create_document("流程.docx", word_path.read_bytes())
        word_chunks = DocumentExtractor.extract(word_document, self.store.data_dir / word_document.stored_path)
        self.assertTrue(any(chunk["section"] == "审批流程" for chunk in word_chunks))

        excel_path = self.root / "预算.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "预算"
        sheet.append(["区域", "金额", "说明"])
        sheet.append(["华东", 500000, "超过阈值"])
        workbook.save(excel_path)
        excel_document = self.create_document("预算.xlsx", excel_path.read_bytes())
        excel_chunks = DocumentExtractor.extract(excel_document, self.store.data_dir / excel_document.stored_path)
        self.assertEqual(excel_chunks[0]["sheet"], "预算")
        self.assertIn("第 2-2 行", excel_chunks[0]["location"])
        self.assertEqual(excel_chunks[0]["cell_range"], "A2:C2")

        pdf_path = self.root / "制度.pdf"
        pdf = pymupdf.open()
        page = pdf.new_page()
        page.insert_text((72, 72), "Risk threshold is 50 percent.")
        pdf.save(pdf_path)
        pdf.close()
        pdf_document = self.create_document("制度.pdf", pdf_path.read_bytes())
        pdf_chunks = DocumentExtractor.extract(pdf_document, self.store.data_dir / pdf_document.stored_path)
        self.assertEqual(pdf_chunks[0]["page"], 1)
        self.assertEqual(pdf_chunks[0]["location"], "第 1 页")

    def test_txt_extractor_decodes_gb18030_and_preserves_folder_display_name(self) -> None:
        document = self.create_document(
            "项目资料/制度/供应商阈值.txt",
            "供应商风险阈值为 50 万元。".encode("gb18030"),
        )
        chunks = DocumentExtractor.extract(document, self.store.data_dir / document.stored_path)

        self.assertEqual(document.original_name, "项目资料/制度/供应商阈值.txt")
        self.assertEqual(chunks[0]["text"], "供应商风险阈值为 50 万元。")

    def test_text_chunking_merges_short_paragraphs_and_ignores_extra_blank_lines(self) -> None:
        paragraphs = [f"第 {index} 条：采购申请需要提供完整的依据、预算和审批记录。" for index in range(1, 19)]
        text = "\n\n\n\n".join(paragraphs)

        chunks = DocumentExtractor._text_chunks("采购制度", text, "正文")

        self.assertLess(len(chunks), len(paragraphs) // 3)
        self.assertIn("第 1 条", chunks[0]["text"])
        self.assertIn("第 18 条", chunks[-1]["text"])
        self.assertTrue(all(len(chunk["text"]) <= DocumentExtractor.max_chunk_size for chunk in chunks))

    def test_text_chunking_keeps_markdown_heading_with_following_content(self) -> None:
        text = "# 报销范围\n\n" + "报销资料必须真实、完整。\n\n" * 8

        chunks = DocumentExtractor._text_chunks("报销制度", text, "正文")

        self.assertTrue(chunks)
        self.assertTrue(chunks[0]["text"].startswith("# 报销范围"))
        self.assertTrue(all(len(chunk["text"]) <= DocumentExtractor.max_chunk_size for chunk in chunks))

    def test_profiles_change_automatic_chunk_granularity_with_bounded_sizes(self) -> None:
        text = "\n\n".join(f"第 {index} 条：" + "审批条件、责任人和例外情况需要完整记录。" * 14 for index in range(1, 7))
        precise = DocumentExtractor._text_chunks(
            "流程", text, "正文", settings=DocumentExtractor.profile_settings("precise")
        )
        context = DocumentExtractor._text_chunks(
            "流程", text, "正文", settings=DocumentExtractor.profile_settings("context")
        )

        self.assertGreater(len(precise), len(context))
        self.assertTrue(all(len(chunk["text"]) <= 700 for chunk in precise))
        self.assertTrue(all(len(chunk["text"]) <= 1200 for chunk in context))

    def test_manual_merge_only_changes_derived_chunks_and_survives_reindex(self) -> None:
        document = self.create_document("流程.txt", b"placeholder")
        original_chunks = [
            {"title": "流程", "text": "事项甲：先核验预算。", "location": "正文"},
            {"title": "流程", "text": "事项甲：再提交审批。", "location": "正文"},
            {"title": "流程", "text": "事项乙：归档材料。", "location": "正文"},
        ]
        self.store.replace_document_chunks(document.id, original_chunks)
        self.store.add_document_chunk_merge(document.id, 0, 1)

        adjusted = self.store.apply_document_chunk_adjustments(self.store.get_document(document.id), original_chunks)
        self.assertEqual(len(adjusted), 2)
        self.assertIn("先核验预算", adjusted[0]["text"])
        self.assertIn("再提交审批", adjusted[0]["text"])
        self.assertEqual((adjusted[0]["source_start_order"], adjusted[0]["source_end_order"]), (0, 1))

        stored = self.store.replace_document_chunks(document.id, adjusted)
        self.assertEqual((stored[0].source_start_order, stored[0].source_end_order), (0, 1))
        self.assertEqual(self.store.document_chunk_adjustment_count(document.id), 1)
        self.store.clear_document_chunk_adjustments(document.id)
        self.assertEqual(self.store.document_chunk_adjustment_count(document.id), 0)

    def test_changing_profile_discards_manual_adjustments(self) -> None:
        document = self.create_document("流程.txt", b"placeholder")
        self.store.replace_document_chunks(document.id, [
            {"text": "第一段", "location": "正文"},
            {"text": "第二段", "location": "正文"},
        ])
        self.store.add_document_chunk_merge(document.id, 0, 1)

        updated = self.store.set_document_chunk_profile(document.id, "precise")

        self.assertEqual(updated.chunk_profile, "precise")
        self.assertEqual(self.store.document_chunk_adjustment_count(document.id), 0)

    def test_replacing_document_keeps_id_and_removes_stale_chunks(self) -> None:
        document = self.create_document("旧版本.txt", "旧规则：阈值为 10 万元。".encode("utf-8"))
        old_path = self.store.data_dir / document.stored_path
        self.store.replace_document_chunks(document.id, [{"text": "旧规则：阈值为 10 万元。", "location": "正文"}])

        replaced = self.store.replace_document_from_data_url(
            document.id,
            "制度/新版本.txt",
            "data:text/plain;base64," + base64.b64encode("新规则：阈值为 50 万元。".encode("utf-8")).decode("ascii"),
        )

        self.assertEqual(replaced.id, document.id)
        self.assertEqual(replaced.original_name, "制度/新版本.txt")
        self.assertEqual(replaced.status, "pending")
        self.assertEqual(replaced.chunk_count, 0)
        self.assertFalse(old_path.exists())
        self.assertTrue((self.store.data_dir / replaced.stored_path).exists())
        self.assertEqual(self.store.list_document_chunks(document.id), [])

    def test_chunk_level_embeddings_retrieve_later_document_content(self) -> None:
        document = self.create_document("长文.txt", "占位".encode("utf-8"))
        chunks = self.store.replace_document_chunks(document.id, [
            {"title": "长文", "text": "项目背景在这里。", "location": "正文第 1 段"},
            {"title": "长文", "text": "深层结论：审批需要双人复核。", "location": "正文第 20 段"},
        ])
        embedding = EmbeddingService(self.store, DocumentEmbeddingEngine())
        embedding._process_document(document.id)
        matches = embedding.semantic_document_search("深层问题", limit=3, min_score=0.5)
        self.assertEqual(matches[0][0].id, chunks[1].id)

        rag = RagPipeline(self.store, embedding)
        retrieved = rag.retrieve("深层结论是什么", limit=3, semantic_min_score=0.0)
        document_result = next(item for item in retrieved if item.get("source_kind") == "document")
        self.assertEqual(document_result["location"], "正文第 20 段")
        source = rag.source_payloads("深层结论是什么", [document_result])[0]
        self.assertEqual(source["type"], "document")
        self.assertEqual(source["location"], "正文第 20 段")
        self.assertEqual(source["chunk_id"], chunks[1].id)
        self.assertEqual(source["chunk_order"], 1)
        self.assertEqual(source["url"], f"/api/documents/{document.id}/file?chunk_id={chunks[1].id}")
        embedding.shutdown()

    def test_chunk_level_embeddings_retrieve_later_note_content(self) -> None:
        note = self.store.create_note(
            "长笔记",
            "项目背景在这里。\n\n" + "前置说明。" * 120 + "\n\n深层结论：预算变更需要双人复核。",
        )
        embedding = EmbeddingService(self.store, DocumentEmbeddingEngine())
        embedding._process(note.id)
        matches = embedding.semantic_note_chunk_search("深层问题", limit=3, min_score=0.5)
        self.assertIn("深层结论", matches[0][0].text)

        rag = RagPipeline(self.store, embedding)
        retrieved = rag.retrieve("深层问题", limit=3, semantic_min_score=0.5)
        self.assertTrue(any("深层结论" in item["raw_text"] for item in retrieved))
        embedding.shutdown()

    def test_delete_document_removes_original_and_derived_index(self) -> None:
        document = self.create_document("删除测试.txt", "可删除的资料".encode("utf-8"))
        stored_path = self.store.data_dir / document.stored_path
        self.store.replace_document_chunks(document.id, [{"text": "可删除的资料", "location": "正文"}])
        self.assertTrue(stored_path.exists())
        self.store.delete_document(document.id)
        self.assertFalse(stored_path.exists())
        self.assertEqual(self.store.document_count(), 0)
        self.assertEqual(self.store.search_document_chunks("删除"), [])

    def test_excel_skips_a_single_cell_title_and_bounds_long_cells(self) -> None:
        excel_path = self.root / "长表.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["项目风险台账"])
        sheet.append(["项目", "风险说明", "责任人"])
        sheet.append(["项目 A", "风险" * 1400, "张三"])
        workbook.save(excel_path)
        document = self.create_document("长表.xlsx", excel_path.read_bytes())

        chunks = DocumentExtractor.extract(document, self.store.data_dir / document.stored_path)

        self.assertTrue(chunks)
        self.assertTrue(all(len(chunk["text"]) <= DocumentExtractor.max_chunk_size for chunk in chunks))
        self.assertTrue(all(chunk["cell_range"].startswith("A3:C3") for chunk in chunks))
