from __future__ import annotations

import re
import tempfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from typing import Callable

from pocket_memory.chunking import split_note_text
from pocket_memory.storage import Document, NoteStore


class DocumentExtractor:
    """Extract locally-readable content while retaining source locations."""

    # Most business-document paragraphs are short. Buffer them into a useful
    # retrieval unit instead of treating every visual paragraph as one chunk.
    target_size = 720
    max_chunk_size = 920
    min_chunk_size = 260
    overlap = 100
    PROFILES = {
        "precise": {"target_size": 500, "max_chunk_size": 700, "min_chunk_size": 180, "overlap": 80},
        "standard": {"target_size": 720, "max_chunk_size": 920, "min_chunk_size": 260, "overlap": 100},
        "context": {"target_size": 900, "max_chunk_size": 1200, "min_chunk_size": 360, "overlap": 120},
    }

    @classmethod
    def profile_settings(cls, profile: str | None) -> dict[str, int]:
        """Return a bounded profile; the standard profile remains the default."""
        return dict(cls.PROFILES.get(str(profile or "").lower(), cls.PROFILES["standard"]))

    @classmethod
    def extract(
        cls,
        document: Document,
        path: Path,
        ocr: Callable[[Path], str] | None = None,
        profile: str | None = None,
    ) -> list[dict]:
        handlers = {
            "txt": cls._extract_txt,
            "docx": cls._extract_docx,
            "xlsx": cls._extract_xlsx,
            "pdf": cls._extract_pdf,
        }
        handler = handlers.get(document.file_type.lower())
        if handler is None:
            raise ValueError(f"不支持的文件类型：{document.file_type}")
        settings = cls.profile_settings(profile or getattr(document, "chunk_profile", "standard"))
        return cls._bound_chunks(handler(document, path, ocr, settings), settings)

    @classmethod
    def _extract_txt(cls, document: Document, path: Path, _ocr=None, settings=None) -> list[dict]:
        data = path.read_bytes()
        for encoding in ("utf-8-sig", "utf-8", "gb18030", "utf-16"):
            try:
                text = data.decode(encoding)
                break
            except UnicodeDecodeError:
                text = ""
        if not text.strip():
            raise ValueError("TXT 文件没有可读取的文字")
        return cls._text_chunks(document.title, text, "正文", settings=settings)

    @classmethod
    def _extract_docx(cls, document: Document, path: Path, _ocr=None, settings=None) -> list[dict]:
        from docx import Document as WordDocument

        word = WordDocument(path)
        chunks: list[dict] = []
        section = ""
        paragraphs: list[str] = []

        def flush_paragraphs() -> None:
            nonlocal paragraphs
            if paragraphs:
                location = f"章节：{section}" if section else "正文"
                chunks.extend(cls._text_chunks(document.title, "\n\n".join(paragraphs), location, section=section, settings=settings))
                paragraphs = []

        for paragraph in word.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style = str(getattr(paragraph.style, "name", ""))
            if style.lower().startswith("heading") or style.startswith("标题"):
                flush_paragraphs()
                section = text
            else:
                paragraphs.append(text)
        flush_paragraphs()

        for table_index, table in enumerate(word.tables, 1):
            rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in table.rows]
            rows = [row for row in rows if any(row)]
            if not rows:
                continue
            headers = rows[0]
            batch: list[str] = []
            start_row = 2
            for row_index, row in enumerate(rows[1:] or rows, 2 if len(rows) > 1 else 1):
                values = [f"{headers[i] if i < len(headers) else f'列{i + 1}'}：{value}" for i, value in enumerate(row) if value]
                line = "；".join(values)
                if batch and len("\n".join(batch + [line])) > settings["target_size"]:
                    end_row = row_index - 1
                    chunks.append(cls._table_chunk(document.title, headers, batch, table_index, start_row, end_row, section))
                    batch = []
                    start_row = row_index
                if line:
                    batch.append(line)
            if batch:
                chunks.append(cls._table_chunk(document.title, headers, batch, table_index, start_row, start_row + len(batch) - 1, section))
        if not chunks:
            raise ValueError("Word 文档没有可提取的正文或表格")
        return chunks

    @classmethod
    def _table_chunk(cls, title, headers, rows, table_index, start_row, end_row, section) -> dict:
        section_label = f"章节：{section}，" if section else ""
        return {
            "title": title,
            "text": "表头：" + " | ".join(headers) + "\n" + "\n".join(rows),
            "location": f"{section_label}表格 {table_index}，第 {start_row}-{end_row} 行",
            "section": section,
            "cell_range": f"表格{table_index}:第{start_row}-{end_row}行",
        }

    @classmethod
    def _extract_xlsx(cls, document: Document, path: Path, _ocr=None, settings=None) -> list[dict]:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=False)
        displayed_workbook = load_workbook(path, read_only=True, data_only=True)
        chunks: list[dict] = []
        try:
            for sheet in workbook.worksheets:
                displayed_sheet = displayed_workbook[sheet.title]
                rows = []
                for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    displayed = [cell.value for cell in displayed_sheet[row_number]]
                    values = [cls._cell_text(value, displayed[index] if index < len(displayed) else None) for index, value in enumerate(row)]
                    if any(values):
                        rows.append((row_number, values))
                if not rows:
                    continue
                header_index = cls._header_row_index(rows)
                header_row, headers = rows[header_index]
                headers = [value or f"列{index + 1}" for index, value in enumerate(headers)]
                batch: list[str] = []
                start_row = None
                end_row = None
                data_rows = rows[header_index + 1:] or rows[header_index:]
                for row_number, values in data_rows:
                    line = "；".join(
                        f"{headers[index] if index < len(headers) else f'列{index + 1}'}：{value}"
                        for index, value in enumerate(values) if value
                    )
                    if not line:
                        continue
                    proposed = "\n".join(batch + [line])
                    if batch and len(proposed) > settings["target_size"]:
                        chunks.append(cls._sheet_chunk(document.title, sheet.title, headers, batch, start_row, end_row))
                        batch = []
                        start_row = None
                    if start_row is None:
                        start_row = row_number
                    batch.append(line)
                    end_row = row_number
                if batch:
                    chunks.append(cls._sheet_chunk(document.title, sheet.title, headers, batch, start_row, end_row))
        finally:
            workbook.close()
            displayed_workbook.close()
        if not chunks:
            raise ValueError("Excel 文件没有可提取的单元格内容")
        return chunks

    @staticmethod
    def _cell_text(value, displayed_value=None) -> str:
        if value is None:
            return ""
        text = str(value).strip().replace("\n", " ")
        if text.startswith("=") and displayed_value not in (None, ""):
            displayed = str(displayed_value).strip().replace("\n", " ")
            if displayed and displayed != text:
                return f"{displayed}（公式：{text}）"
        return text

    @staticmethod
    def _header_row_index(rows: list[tuple[int, list[str]]]) -> int:
        """Skip single-cell sheet titles and choose the first plausible header."""
        for index, (_, values) in enumerate(rows[:10]):
            populated = [value for value in values if value]
            text_cells = [value for value in populated if not re.fullmatch(r"[-+]?\d+(?:\.\d+)?", value)]
            if len(populated) >= 2 and len(text_cells) >= 2:
                return index
        return 0

    @classmethod
    def _sheet_chunk(cls, title, sheet, headers, rows, start_row, end_row) -> dict:
        last_column = max((index + 1 for index, header in enumerate(headers) if header), default=1)
        column_label = cls._column_label(last_column)
        return {
            "title": title,
            "text": f"工作表：{sheet}\n表头：" + " | ".join(headers) + "\n" + "\n".join(rows),
            "location": f"工作表：{sheet}，第 {start_row}-{end_row} 行",
            "sheet": sheet,
            "cell_range": f"A{start_row}:{column_label}{end_row}",
        }

    @staticmethod
    def _column_label(number: int) -> str:
        label = ""
        while number:
            number, remainder = divmod(number - 1, 26)
            label = chr(65 + remainder) + label
        return label or "A"

    @classmethod
    def _bound_chunks(cls, chunks: list[dict], settings: dict[str, int] | None = None) -> list[dict]:
        """Enforce the documented hard ceiling for table rows and long cells."""
        settings = settings or cls.profile_settings("standard")
        bounded: list[dict] = []
        for chunk in chunks:
            text = str(chunk.get("text", ""))
            if len(text) <= settings["max_chunk_size"]:
                bounded.append(chunk)
                continue
            lines = text.splitlines()
            header = lines[0] if lines and lines[0].startswith("表头：") else ""
            body = "\n".join(lines[1:] if header else lines)
            limit = max(120, settings["max_chunk_size"] - len(header) - (1 if header else 0))
            pieces = split_note_text(body, target_size=limit, overlap=0) or [body[:limit]]
            for index, piece in enumerate(pieces, 1):
                item = dict(chunk)
                item["text"] = f"{header}\n{piece}".strip() if header else piece
                item["location"] = f"{chunk.get('location', '正文')}，片段 {index}/{len(pieces)}"
                bounded.append(item)
        return bounded

    @classmethod
    def _extract_pdf(cls, document: Document, path: Path, ocr=None, settings=None) -> list[dict]:
        import pymupdf

        pdf = pymupdf.open(path)
        chunks: list[dict] = []
        try:
            with tempfile.TemporaryDirectory(prefix="pocket-memory-pdf-") as temporary:
                temporary_dir = Path(temporary)
                for page_number, page in enumerate(pdf, 1):
                    text = page.get_text("text").strip()
                    location = f"第 {page_number} 页"
                    if not text and ocr is not None:
                        image_path = temporary_dir / f"page-{page_number}.png"
                        pixmap = page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5), alpha=False)
                        pixmap.save(image_path)
                        text = str(ocr(image_path) or "").strip()
                        location += "（OCR）"
                    if text:
                        chunks.extend(cls._text_chunks(document.title, text, location, page=page_number, settings=settings))
        finally:
            pdf.close()
        if not chunks:
            raise ValueError("PDF 没有可提取文字；扫描件需要安装并启用 OCR 模型")
        return chunks

    @classmethod
    def _text_chunks(cls, title: str, text: str, location: str, settings: dict[str, int] | None = None, **metadata) -> list[dict]:
        settings = settings or cls.profile_settings("standard")
        text = re.sub(r"\r\n?", "\n", str(text or "").strip())
        text = re.sub(r"\n[ \t]*\n(?:[ \t]*\n)+", "\n\n", text)
        if not text:
            return []
        result: list[dict] = []
        current: list[str] = []
        current_heading = ""
        last_chunk_heading = ""

        def current_text(parts: list[str]) -> str:
            prefix = f"{current_heading}\n" if current_heading else ""
            return (prefix + "\n\n".join(parts)).strip()

        def flush() -> None:
            nonlocal current, last_chunk_heading
            if not current:
                return
            chunk_text = current_text(current)
            if result and current_heading == last_chunk_heading and len(chunk_text) < settings["min_chunk_size"]:
                merged = f"{result[-1]['text']}\n\n{chunk_text}".strip()
                if len(merged) <= settings["max_chunk_size"]:
                    result[-1]["text"] = merged
                    current = []
                    return
            result.append({"title": title, "text": chunk_text, "location": location, **metadata})
            last_chunk_heading = current_heading
            current = []

        def split_long_paragraph(paragraph: str, limit: int) -> list[str]:
            sentences = [item.strip() for item in re.split(r"(?<=[。！？；;.!?])\s*|\n+", paragraph) if item.strip()]
            if not sentences:
                sentences = [paragraph]
            pieces: list[str] = []
            piece = ""
            for sentence in sentences:
                while len(sentence) > limit:
                    if piece:
                        pieces.append(piece)
                        piece = ""
                    pieces.append(sentence[:limit])
                    sentence = sentence[limit - settings["overlap"]:]
                candidate = f"{piece}\n{sentence}".strip() if piece else sentence
                if piece and len(candidate) > limit:
                    pieces.append(piece)
                    piece = sentence
                else:
                    piece = candidate
            if piece:
                pieces.append(piece)
            return pieces

        parts = [part.strip() for part in re.split(r"\n\s*\n", text) if part.strip()]
        for part in parts:
            if re.fullmatch(r"#{1,6}\s+.+", part):
                flush()
                current_heading = part
                continue
            # Keep room for a Markdown heading when it is repeated with its
            # content, so the absolute chunk limit remains meaningful.
            unit_limit = max(settings["min_chunk_size"], settings["max_chunk_size"] - len(current_heading) - 1)
            for unit in split_long_paragraph(part, unit_limit):
                candidate = current_text(current + [unit])
                if current and len(candidate) > settings["target_size"]:
                    # A short near-target chunk is more useful than a tiny one;
                    # never exceed the hard limit merely to avoid a split.
                    if len(current_text(current)) < settings["min_chunk_size"] and len(candidate) <= settings["max_chunk_size"]:
                        current.append(unit)
                    else:
                        flush()
                        current.append(unit)
                else:
                    current.append(unit)
        flush()
        return result


class DocumentImportService:
    """Asynchronously extract files, then hand their small chunks to embedding."""

    def __init__(self, store: NoteStore, embedding_service, ocr_engine=None) -> None:
        self.store = store
        self.embedding_service = embedding_service
        self.ocr_engine = ocr_engine
        self.executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix="pocket-memory-document")

    def enqueue(self, document_id: int) -> None:
        self.store.set_document_status(document_id, "extracting")
        self.executor.submit(self._process, document_id)

    def reindex(self, document_id: int) -> None:
        self.enqueue(document_id)

    def status(self) -> dict:
        documents = self.store.list_documents()
        return {
            "total": len(documents),
            "ready": sum(document.status == "ready" for document in documents),
            "processing": sum(document.status in {"pending", "extracting", "indexing"} for document in documents),
            "failed": sum(document.status == "failed" for document in documents),
        }

    def shutdown(self) -> None:
        self.executor.shutdown(wait=False, cancel_futures=True)

    def _process(self, document_id: int) -> None:
        document = None
        try:
            document = self.store.get_document(document_id)
            path = self.store.data_dir / document.stored_path
            ocr = self._recognize if self.ocr_engine is not None and self.ocr_engine.available else None
            chunks = DocumentExtractor.extract(document, path, ocr, profile=document.chunk_profile)
            chunks = self.store.apply_document_chunk_adjustments(document, chunks)
            if not chunks:
                raise ValueError("没有可索引的文本")
            replaced = self.store.replace_document_chunks(document_id, chunks, expected_sha256=document.sha256)
            if replaced is not None:
                self.embedding_service.enqueue_document(document_id)
        except KeyError:
            return
        except Exception as exc:
            if document is not None and self.store.document_matches_sha256(document_id, document.sha256):
                self.store.set_document_status(document_id, "failed", f"{type(exc).__name__}: {exc}")

    def _recognize(self, path: Path) -> str:
        return self.ocr_engine.recognize(path)
